"""Sinh 3 file sao kê giả cho test import (Phase 14) — dữ liệu bịa, không phải sao kê thật.

    (cd backend && python tests/fixtures/imports/generate_fixtures.py)

Cùng 25 dòng logic: 22 hợp lệ, 2 dòng trùng (cùng ngày / số tiền / chiều / nội dung với một dòng
trước), 1 dòng lỗi (ngày 32/13/2026).

- `vcb_statement.csv`: UTF-8 có BOM, dấu `;`, cột ghi nợ / ghi có riêng, có dòng mở đầu.
- `techcombank_statement.csv`: windows-1258 (cp1258, dấu thanh tổ hợp), dấu `,`, một cột số tiền
  có dấu (âm = chi).
- `statement.xlsx`: ô ngày kiểu date, ô số kiểu số, có dòng mở đầu.
"""

import csv
import io
import unicodedata
from datetime import date
from pathlib import Path

HERE = Path(__file__).parent

# (ngày, nội dung, ghi nợ, ghi có)
ROWS: list[tuple[str, str, int, int]] = [
    ("01/09/2026", "Lương tháng 08/2026 công ty ABC", 0, 25_000_000),
    ("02/09/2026", "GRAB*TRIP 1234 HCM", 45_000, 0),
    ("02/09/2026", "Ăn trưa quán bún bò", 65_000, 0),
    ("03/09/2026", "THANH TOAN TIEN DIEN EVN THANG 8", 1_500_000, 0),
    ("04/09/2026", "GRAB*TRIP 5678 HCM", 52_000, 0),
    ("05/09/2026", "BACH HOA XANH Q7", 320_000, 0),
    ("05/09/2026", "HIGHLANDS COFFEE VINCOM", 59_000, 0),
    ("05/09/2026", "HIGHLANDS COFFEE VINCOM", 59_000, 0),  # trùng dòng trên
    ("06/09/2026", "SHOPEE PAY DON HANG 260906", 450_000, 0),
    ("06/09/2026", "SHOPEE PAY DON HANG 260906", 450_000, 0),  # trùng dòng trên
    ("07/09/2026", "NETFLIX.COM", 260_000, 0),
    ("08/09/2026", "GRAB*TRIP 9012 HCM", 38_000, 0),
    ("09/09/2026", "CK tiền nhà tháng 9 cho chủ nhà", 5_000_000, 0),
    ("10/09/2026", "NHA THUOC LONG CHAU", 180_000, 0),
    ("11/09/2026", "VIETTEL CUOC INTERNET T9", 230_000, 0),
    ("32/13/2026", "PHI DICH VU SMS BANKING", 11_000, 0),  # ngày lỗi
    ("12/09/2026", "Thanh toán dự án thiết kế logo", 0, 3_000_000),
    ("13/09/2026", "CGV CINEMAS VINCOM", 120_000, 0),
    ("14/09/2026", "Sửa xe máy Honda", 2_800_000, 0),
    ("15/09/2026", "TRA SUA PHUC LONG", 85_000, 0),
    ("16/09/2026", "WINMART THAO DIEN", 410_000, 0),
    ("18/09/2026", "Học phí tiếng Anh tháng 9", 1_200_000, 0),
    ("20/09/2026", "HOAN TIEN DON HANG LAZADA", 0, 500_000),
    ("22/09/2026", "BE CAR CHUYEN DI", 95_000, 0),
    ("25/09/2026", "WINMART THAO DIEN", 275_000, 0),
]

OPENING_BALANCE = 12_000_000
_TONES = {"\u0300", "\u0301", "\u0303", "\u0309", "\u0323"}


def vnd(amount: int) -> str:
    """25000000 → "25.000.000"; 0 → ô trống."""
    return f"{amount:,}".replace(",", ".") if amount else ""


def balances() -> list[int]:
    result, balance = [], OPENING_BALANCE
    for _, _, debit, credit in ROWS:
        balance += credit - debit
        result.append(balance)
    return result


def write_vcb() -> None:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow(["SAO KÊ TÀI KHOẢN (DỮ LIỆU GIẢ)"])
    writer.writerow(["Số tài khoản: 0001234567890", "Từ ngày 01/09/2026 đến 30/09/2026"])
    writer.writerow([])
    writer.writerow(
        [
            "STT",
            "Ngày giao dịch",
            "Số tham chiếu",
            "Số tiền ghi nợ",
            "Số tiền ghi có",
            "Số dư",
            "Nội dung chi tiết",
        ]
    )
    for index, ((day, note, debit, credit), balance) in enumerate(
        zip(ROWS, balances(), strict=True), 1
    ):
        reference = f"FT26{index:03d}88{index:04d}"
        writer.writerow([index, day, reference, vnd(debit), vnd(credit), vnd(balance), note])
    (HERE / "vcb_statement.csv").write_bytes(b"\xef\xbb\xbf" + out.getvalue().encode("utf-8"))


def encode_cp1258(text: str) -> bytes:
    """cp1258 chỉ có sẵn chữ gốc + dấu mũ / trăng / móc; dấu thanh là ký tự tổ hợp riêng."""
    out = bytearray()
    for ch in unicodedata.normalize("NFC", text):
        try:
            out += ch.encode("cp1258")
            continue
        except UnicodeEncodeError:
            pass
        decomposed = unicodedata.normalize("NFD", ch)
        base = decomposed[0] + "".join(m for m in decomposed[1:] if m not in _TONES)
        tones = "".join(m for m in decomposed[1:] if m in _TONES)
        out += unicodedata.normalize("NFC", base).encode("cp1258") + tones.encode("cp1258")
    return bytes(out)


def write_techcombank() -> None:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=",", lineterminator="\r\n")
    writer.writerow(["Ngày", "Diễn giải", "Số tiền", "Số dư"])
    for (day, note, debit, credit), balance in zip(ROWS, balances(), strict=True):
        amount = f"-{vnd(debit)}" if debit else vnd(credit)
        writer.writerow([day, note, amount, vnd(balance)])
    (HERE / "techcombank_statement.csv").write_bytes(encode_cp1258(out.getvalue()))


def write_xlsx() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sao ke"
    sheet.append(["SAO KÊ TÀI KHOẢN THANH TOÁN (DỮ LIỆU GIẢ)"])
    sheet.append(["Chủ tài khoản: NGUYEN VAN A"])
    sheet.append(["Ngày GD", "Mô tả", "Ghi nợ", "Ghi có", "Số dư"])
    for (day, note, debit, credit), balance in zip(ROWS, balances(), strict=True):
        d, m, y = (int(part) for part in day.split("/"))
        try:
            cell_date: date | str = date(y, m, d)
        except ValueError:
            cell_date = day
        sheet.append([cell_date, note, debit or None, credit or None, balance])
        if isinstance(cell_date, date):
            sheet.cell(row=sheet.max_row, column=1).number_format = "DD/MM/YYYY"
    workbook.properties.creator = "pfa-fixtures"
    workbook.save(HERE / "statement.xlsx")


if __name__ == "__main__":
    write_vcb()
    write_techcombank()
    write_xlsx()
    print("Đã sinh:", ", ".join(p.name for p in sorted(HERE.iterdir()) if p.suffix != ".py"))
